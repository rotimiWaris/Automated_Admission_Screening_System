"""Generate clearly fictional applicants in the supplied 21-column layout."""
from datetime import datetime
from pathlib import Path
import shutil
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))
from screening.applicant_data import HEADERS


def generate(destination, twenty=False):
    wb = Workbook()
    wb.properties.title = 'SYNTHETIC TEST DATA — MAPOLY student register'
    wb.properties.description = 'Fictional applicants for software testing. Not real MAPOLY student records.'
    sheet = wb.active
    sheet.title = 'Synthetic applicants'
    sheet.append(HEADERS)
    cases = [
        ('Accountancy', 210, 'English Language; Mathematics; Economics; Commerce', 'Economics;Commerce;Financial Accounting', 'Qualified'),
        ('Computer Engineering', 99, 'English Language; Mathematics; Physics; Chemistry', 'Physics;Chemistry;Computer Studies', 'Eligible'),
        ('Civil Engineering', 180, 'English Language; Mathematics; Physics; Chemistry', 'Physics;Chemistry;Building Construction', 'Unqualified'),
        ('Computer Science', 200, 'English Language; Mathematics; Physics; Chemistry', 'Physics;Chemistry;Biology', ''),
        ('Mass Communication', 170, 'English Language; Government; Literature in English; Economics', 'Literature in English;Government;History', 'Qualified'),
        ('Library and Information Science', 190, 'English Language; History; Music; French', 'History;Music;French', ''),
        ('Science Laboratory Technology', 220, 'English Language; Mathematics; Biology; Physics', 'Physics;Chemistry;Economics;Biology', 'Eligible'),
        ('Urban and Regional Planning', 180, 'English Language; Geography; Physics; Chemistry', 'Physics;Chemistry;Biology', ''),
        ('Electrical/Electronics Engineering', 230, 'English Language; Mathematics; Physics; Chemistry', 'Physics;Chemistry;Technical Drawing', 'Qualified'),
        ('Food Technology', 210, 'English Language; Biology; Chemistry; Economics', 'Chemistry;Biology;Economics;Physics', ''),
        ('Computer Engineering', '', 'English Language; Mathematics; Physics; Chemistry', 'Physics;Chemistry;Computer Studies', 'Eligible'),
        ('Accountancy', 180, 'English Language; Mathematics; Economics;Commerce', 'Economics;Commerce;Financial Accounting', 'Unqualified'),
    ]
    if twenty:
        cases += cases[:8]
    for number, (programme, score, utme, electives, remark) in enumerate(cases, 1):
        subjects = ['English Language', 'Mathematics'] + electives.split(';')
        grades = [subject + ': C6' for subject in subjects]
        if twenty:
            grade_cycle = ['A1', 'B2', 'B3', 'C4', 'C5', 'C6']
            grades = [f'{subject}: {grade_cycle[(number + index - 1) % len(grade_cycle)]}' for index, subject in enumerate(subjects)]
        if number == 5:
            grades[0] = 'English Language: F9'
        if number == 7:
            grades[-1] = 'Biology: E8'
        if number == 10:
            grades[-1] = 'Physics: F9'
        if twenty and number in [15, 18]:
            grades[1] = 'Mathematics: D7'
        if twenty and number == 19:
            grades[-1] = 'Biology: E8'
        first = 'WAEC School 2025; ' + '; '.join(grades)
        second = ''
        if number == 9 or (twenty and number in [13, 16, 20]):
            first = 'WAEC School 2025; ' + '; '.join(grades[:3])
            second = 'NECO School 2026; ' + '; '.join(grades[3:])
        age = 22 if number == 6 else None
        dob = None if number == 6 else datetime(2005, (number % 12) + 1, 15)
        prefix = 'TEST2' if twenty else 'TEST'
        values = [number, f'{prefix}-FORM-{number:03}', f'{prefix}-JAMB-{number:03}', f'{prefix} Applicant {number:02}',
                  prefix, 'Applicant', f'{number:02}', f'{prefix.lower()}applicant{number:02}@example.invalid', f'0000000{number:04}',
                  dob, age, 'Moshood Abiola Polytechnic', programme, utme, score, 2 if second else 1,
                  first, second, 'Uploaded', datetime(2026, 9, 20, 10, 0), remark]
        sheet.append(values)
        sheet.cell(number+1, 10).number_format = 'yyyy-mm-dd'
        sheet.cell(number+1, 20).number_format = 'yyyy-mm-dd hh:mm'
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='175B3A')
        cell.font = Font(bold=True, color='FFFFFF')
    sheet['A1'].comment = Comment('SYNTHETIC TEST DATA. Every applicant in this workbook is fictional.', 'Project')
    sheet['J1'].comment = Comment('Use an Excel date, YYYY-MM-DD or DD/MM/YYYY. A missing age is calculated on the processing date.', 'Project')
    sheet['Q1'].comment = Comment('Format: WAEC School 2025; English Language: C6; Mathematics: B3. Use the same format for sitting 2. The real school export may need additional formatting support.', 'Project')
    sheet['U1'].comment = Comment('Existing remarks are optional. The software preserves and compares them with its independently calculated remark.', 'Project')
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 22
    for column in ['N','Q','R']:
        sheet.column_dimensions[column].width = 65
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination)
    return len(cases)


if __name__ == '__main__':
    twenty = '--twenty' in sys.argv
    filename = 'mapoly-student-register-20-students.xlsx' if twenty else 'mapoly-student-register-sample.xlsx'
    path = ROOT / 'samples/synthetic-admissions/tabular' / filename
    count = generate(path, twenty=twenty)
    public = ROOT / 'frontend/public/templates' / filename
    public.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(path, public)
    print(f'Created {count} fictional applicants, 21 columns: {path}')
