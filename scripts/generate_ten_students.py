"""Create another fictional register with fresh applicant identifiers."""
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
folder = root / 'samples/synthetic-admissions/tabular'
workbook = load_workbook(folder / 'mapoly-student-register-20-students.xlsx')
sheet = workbook.active
sheet.delete_rows(12, sheet.max_row - 11)
for row in range(2, 12):
    number = row - 1
    for column, value in {
        2: f'TEST3-FORM-{number:03}', 3: f'TEST3-JAMB-{number:03}',
        4: f'TEST3 Applicant {number:02}', 5: 'TEST3',
        8: f'test3applicant{number:02}@example.invalid',
    }.items():
        sheet.cell(row, column).value = value
sheet.auto_filter.ref = sheet.dimensions
destination = folder / 'mapoly-student-register-10-students.xlsx'
workbook.save(destination)
verified = load_workbook(destination).active
assert verified.max_row == 11 and verified.max_column == 21
assert len({verified.cell(row, 3).value for row in range(2, 12)}) == 10
print(destination)
