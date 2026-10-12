"""Reusable CSV/XLSX readers. No OCR, student schema or formula evaluation."""
import csv
from dataclasses import dataclass
from io import BytesIO, StringIO
from itertools import islice
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from openpyxl import load_workbook
from defusedxml.common import DefusedXmlException
from .service import OCRError


@dataclass(frozen=True)
class TableRow:
    sheet: str
    number: int
    values: dict[str, str]


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_table(raw: bytes, filename: str, max_rows: int = 1000) -> list[TableRow]:
    if not raw or len(raw) > 10 * 1024 * 1024:
        raise OCRError('Table is empty or exceeds 10 MB.')
    rows = []
    def collect(source, sheet):
        headers = None
        for number, cells in enumerate(source, 1):
            if number > max_rows + 20:
                raise OCRError('Too many worksheet rows, including blank rows.')
            values = [cell_text(value) for value in cells]
            if any(len(value) > 2000 for value in values) or len(values) > 80:
                raise OCRError('Table exceeds the 80-column or 2,000-character cell limit.')
            if not any(values):
                continue
            if headers is None:
                while values and not values[-1]:
                    values.pop()
                if not all(values) or len({v.casefold() for v in values}) != len(values):
                    raise OCRError(f'{sheet}: column headers must be nonempty and unique.')
                headers = values
                continue
            if any(values[len(headers):]):
                raise OCRError(f'{sheet} row {number}: values exist without column headers.')
            values += [''] * max(0, len(headers) - len(values))
            rows.append(TableRow(sheet, number, dict(zip(headers, values))))
            if len(rows) > max_rows:
                raise OCRError(f'Table exceeds the {max_rows}-row limit.')
    suffix = filename.lower().rsplit('.', 1)[-1]
    if suffix == 'csv':
        try:
            text = raw.decode('utf-8-sig')
            if '\x00' in text:
                raise OCRError('Save this CSV with UTF-8 encoding.')
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            collect(csv.reader(StringIO(text, newline=''), dialect, strict=True), 'CSV')
        except (UnicodeDecodeError, csv.Error) as exc:
            raise OCRError('Cannot read CSV. Use a valid UTF-8 CSV with a header row.') from exc
    elif suffix == 'xlsx':
        workbook = None
        try:
            with ZipFile(BytesIO(raw)) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(item.file_size for item in entries) > 30 * 1024 * 1024:
                    raise OCRError('Excel workbook is too large when decompressed.')
                if any('vbaproject' in item.filename.lower() for item in entries):
                    raise OCRError('Macro-enabled workbooks are not supported. Save as plain XLSX.')
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            if len(workbook.worksheets) > 10:
                raise OCRError('Excel workbooks may contain at most 10 worksheets.')
            for sheet in workbook.worksheets:
                if sheet.max_column and sheet.max_column > 80:
                    raise OCRError('Excel worksheets may contain at most 80 columns.')
                # Do not trust a workbook's declared dimensions to bound work.
                sheet.reset_dimensions()
                def values():
                    for row in islice(sheet.iter_rows(max_col=81), max_rows + 21):
                        if any(cell.data_type == 'f' for cell in row):
                            raise OCRError(f'{sheet.title}: replace formulas with their values before importing.')
                        result = [cell.value for cell in row]
                        while result and result[-1] is None:
                            result.pop()
                        yield result
                collect(values(), sheet.title)
        except OCRError:
            raise
        except (BadZipFile, OSError, ValueError, KeyError, RuntimeError, ParseError, DefusedXmlException) as exc:
            raise OCRError('Cannot open this Excel workbook. Use a valid, unencrypted XLSX file.') from exc
        finally:
            if workbook is not None:
                workbook.close()
    else:
        raise OCRError('Use CSV or Excel .xlsx. Convert older .xls files to .xlsx first.')
    if not rows:
        raise OCRError('The table contains no data rows.')
    return rows
