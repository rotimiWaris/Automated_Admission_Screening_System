import unittest
from io import BytesIO
from openpyxl import Workbook
from result_ocr.tables import read_table
from result_ocr import OCRError


class TableTests(unittest.TestCase):
    def test_csv_preserves_identifiers_and_quoted_fields(self):
        rows = read_table(b'\xef\xbb\xbfName;Registration\r\n"Student, One";00123AB\r\n', 'data.csv')
        self.assertEqual(rows[0].values, {'Name': 'Student, One', 'Registration': '00123AB'})
        self.assertEqual(rows[0].number, 2)

    def test_all_excel_sheets_read_without_recognition(self):
        workbook = Workbook()
        workbook.active.append(['Name', 'Score'])
        workbook.active.append(['First', 245])
        other = workbook.create_sheet('Second')
        other.append(['Name', 'Score'])
        other.append(['Second', 180])
        data = BytesIO(); workbook.save(data); workbook.close()
        rows = read_table(data.getvalue(), 'data.xlsx')
        self.assertEqual([r.values['Score'] for r in rows], ['245', '180'])
        self.assertEqual(rows[1].sheet, 'Second')

    def test_formulas_and_invalid_tables_rejected(self):
        workbook = Workbook(); workbook.active.append(['Score']); workbook.active.append(['=100+100'])
        data = BytesIO(); workbook.save(data); workbook.close()
        for raw, name in [(data.getvalue(), 'data.xlsx'), (b'a,a\n1,2', 'data.csv'), (b'a\n1,2', 'data.csv'), (b'a\n1', 'data.xls')]:
            with self.subTest(name=name), self.assertRaises(OCRError):
                read_table(raw, name)

    def test_row_limit_does_not_silently_truncate(self):
        with self.assertRaises(OCRError):
            read_table(b'name\none\ntwo\n', 'data.csv', max_rows=1)
