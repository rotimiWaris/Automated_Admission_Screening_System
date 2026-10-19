import base64
import json
from unittest.mock import patch
from datetime import date
from io import BytesIO
from pathlib import Path
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core import signing
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from django.db import OperationalError
from openpyxl import load_workbook
from .applicant_data import HEADERS, applicant_metadata, remark_comparison
from .table_import import import_students
from .models import RuleSet, Verification


class ApplicantTests(SimpleTestCase):
    def test_age_at_birthday_and_provided_age(self):
        self.assertEqual(applicant_metadata({'dob': '2005-09-28'}, date(2026, 9, 27))['age'], 20)
        self.assertEqual(applicant_metadata({'dob': '28/09/2005'}, date(2026, 9, 28))['age'], 21)
        provided = applicant_metadata({'dob': '2005-09-28', 'age': '22'}, date(2026, 9, 28))
        self.assertEqual(provided['age'], 22)
        self.assertTrue(provided['warnings'])
        self.assertIsNone(applicant_metadata({'dob': 'bad'})['age'])
        self.assertIsNone(applicant_metadata({'dob': '2099-01-01'})['age'])

    def test_remarks_never_override_decision(self):
        self.assertEqual(remark_comparison('Eligible', 'Not Eligible')['mark'], '×')
        self.assertEqual(remark_comparison('Unqualified', 'Not Eligible')['mark'], '✓')
        self.assertEqual(remark_comparison('Eligible', 'Pending Review')['mark'], '?')
        self.assertEqual(remark_comparison('Unreadable', 'Eligible')['mark'], '?')


class RegisterTests(TestCase):
    def test_thousand_student_import_screen_save_and_export(self):
        from openpyxl import Workbook
        from result_ocr import OCRError
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(HEADERS)
        source = self.drafts[0]['source_record']
        for index in range(1000):
            row = dict(source, **{'Jamb reg': f'BULK-{index:04}', 'Form No.': f'FORM-{index:04}'})
            sheet.append([row[h] for h in HEADERS])
        raw = BytesIO()
        workbook.save(raw)
        self.drafts, _ = import_students(raw.getvalue(), 'bulk.xlsx')
        self.assertEqual(len(self.drafts), 1000)
        response = self.post_batch()
        self.assertEqual(response.status_code, 200, response.content[:300])
        self.assertEqual(response.json()['summary'], {'Eligible': 1000})
        self.assertEqual(Verification.objects.count(), 1000)
        report = load_workbook(BytesIO(base64.b64decode(response.json()['excel']['base64'])))
        self.assertEqual(report.active.max_row, 1001)
        self.assertEqual(report.active.max_column, 21)
        from .mapoly_import import import_register
        from result_ocr.tables import TableRow
        self.drafts = [self.drafts[0]] * 10001
        self.assertEqual(self.post_batch().status_code, 400)
        with self.assertRaisesMessage(OCRError, 'at most 10000 students'):
            import_register([TableRow('School',2,source)] * 10001, 'too-many.xlsx')

    @classmethod
    def setUpTestData(cls):
        call_command('load_compendium', verbosity=0)

    def setUp(self):
        path = Path(settings.BASE_DIR).parent / 'samples/synthetic-admissions/tabular/mapoly-student-register-sample.xlsx'
        self.drafts, _ = import_students(path.read_bytes(), path.name)
        self.user = get_user_model().objects.create_user(username='register-admin')
        self.client.force_login(self.user)

    def post_batch(self, admission_type='ND Full-Time', token_type=None):
        token = signing.dumps({'document': 'register', 'page': 1, 'admission_type': token_type or admission_type, 'drafts': self.drafts}, salt='ocr-batch')
        mappings = {r.programme: {'id': r.id, 'revision': r.revision} for r in RuleSet.objects.filter(active=True)}
        return self.client.post('/api/screen-batch/', json.dumps({'pages': [token], 'course_rules': mappings, 'admission_type': admission_type, 'save': True}), content_type='application/json')

    def test_sample_screening_export_history_and_duplicate_rejection(self):
        self.assertEqual(len(self.drafts), 12)
        self.assertEqual(set(self.drafts[0]['source_record']), set(HEADERS))
        response = self.post_batch()
        self.assertEqual(response.status_code, 200, response.content[:500])
        data = response.json()
        self.assertEqual([r['status'] for r in data['results']], ['Eligible', 'Not Eligible', 'Eligible', 'Eligible', 'Not Eligible', 'Pending Review', 'Eligible', 'Eligible', 'Eligible', 'Not Eligible', 'Pending Review', 'Eligible'])
        self.assertEqual(Verification.objects.count(), 12)
        first = data['results'][0]['result']
        self.assertIsInstance(first['applicant']['age'], int)
        self.assertEqual(first['remark_check']['mark'], '✓')
        self.assertTrue(first['subject_checks'])
        self.assertEqual(Verification.objects.get(pk=first['id']).result['remark'], 'Qualified')
        workbook = load_workbook(BytesIO(base64.b64decode(data['excel']['base64'])))
        self.assertEqual(workbook.sheetnames, ['Screenings'])
        rows = list(workbook['Screenings'].values)
        self.assertEqual(list(rows[0]), ['#', 'FORM NO', 'Jamb reg', 'NAME ON JAMB', 'SURNAME', 'FIRST NAME', 'OTHER NAME', 'E-mail', 'Phone', 'DOB', 'AGE', 'First choice', 'Jamb Programme', 'Jamb Subject', 'Jamb Score', 'Sitting(s)', 'OLevel Result Sitting 1', 'OLevel Result Sitting 2', 'JAMB OLEVEL STATUS', 'Submit Date', 'Remark'])
        self.assertTrue(all(cell.font.bold for cell in workbook['Screenings'][1]))
        self.assertEqual(workbook['Screenings'].max_column, 21)
        self.assertEqual(rows[2][20], 'Unqualified')
        self.assertEqual(rows[1][20], 'Qualified')
        self.assertIn('Previous remark: Eligible', workbook['Screenings']['U3'].comment.text)
        self.assertIn('Previous remark disagrees', workbook['Screenings']['U3'].comment.text)
        self.assertIn('English Language [✓]', rows[1][13])
        self.assertIn('[✓]', rows[9][16])
        self.assertIn('[✓]', rows[9][17])
        self.assertEqual(rows[11][20], 'Pending Review')
        self.assertIn('English Language: C6 [✓]', rows[1][16])
        self.assertIn('English Language: F9', rows[5][16])
        self.assertNotIn('[', rows[5][16])
        self.assertEqual(str(rows[2][14]), '99')
        self.assertIn('JAMB score', workbook['Screenings']['U3'].comment.text)
        repeated = self.post_batch().json()
        self.assertEqual(repeated['summary'], data['summary'])
        self.assertTrue(all(item['save_skipped'] for item in repeated['results']))
        repeat_sheet = load_workbook(BytesIO(base64.b64decode(repeated['excel']['base64'])))['Screenings']
        self.assertEqual(repeat_sheet['U2'].value, 'Qualified')
        self.assertEqual(repeat_sheet['U3'].value, 'Unqualified')
        self.assertIn('[✓]', repeat_sheet['Q2'].value)
        self.assertIn('[✓]', repeat_sheet['N2'].value)
        for row in repeat_sheet:
            for cell in row:
                self.assertNotIn('duplicate', str(cell.value).lower())
                if cell.comment:
                    self.assertNotIn('duplicate', cell.comment.text.lower())
        self.assertEqual(Verification.objects.count(), 12)

    def test_other_admission_types_do_not_use_nd_full_time_rules(self):
        response = self.post_batch('HND Full-Time')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['summary'], {'Pending Review': 12})
        self.assertFalse(Verification.objects.exists())
        self.assertEqual(self.post_batch('ND Part-Time', token_type='ND Full-Time').status_code, 400)

    def test_repeat_upload_uses_current_results_and_marks_incorrect_jamb_subject(self):
        self.post_batch()
        self.drafts[0]['jamb_subjects'] = ['English Language', 'Mathematics', 'French', 'Commerce']
        data = self.post_batch().json()
        sheet = load_workbook(BytesIO(base64.b64decode(data['excel']['base64'])))['Screenings']
        self.assertIn('French', sheet['N2'].value)
        self.assertNotIn('[', sheet['N2'].value)
        self.assertEqual(sheet['U2'].value, 'Unqualified')
        self.assertIn('JAMB subject combination', sheet['U2'].comment.text)
        self.assertEqual(Verification.objects.count(), 12)
        self.assertEqual(Verification.objects.order_by('id').first().result['remark'], 'Qualified')

    def test_upload_endpoint_signs_register_with_selected_route(self):
        path = Path(settings.BASE_DIR).parent / 'samples/synthetic-admissions/tabular/mapoly-student-register-sample.xlsx'
        response = self.client.post('/api/ocr/', {'file': SimpleUploadedFile(path.name, path.read_bytes()), 'admission_type': 'ND Part-Time'})
        self.assertEqual(response.status_code, 422)
        self.assertIn('ND Part-Time register', response.json()['error'])

    def test_export_formula_injection_is_plain_text(self):
        self.drafts[0]['source_record']['Surname'] = '=HYPERLINK("bad")'
        data = self.post_batch().json()
        workbook = load_workbook(BytesIO(base64.b64decode(data['excel']['base64'])))
        self.assertEqual(workbook['Screenings']['E2'].data_type, 's')

    def test_database_lock_rolls_back_batch_and_returns_retry_response(self):
        create = Verification.objects.get_or_create
        calls = 0
        def locked(**kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise OperationalError('database is locked')
            return create(**kwargs)
        with patch('screening.batch_views.Verification.objects.get_or_create', side_effect=locked):
            response = self.post_batch()
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response['Retry-After'], '3')
        self.assertIn('database is busy', response.json()['error'])
        self.assertFalse(Verification.objects.exists())
        self.assertEqual(self.post_batch().status_code, 200)
        self.assertEqual(Verification.objects.count(), 12)
