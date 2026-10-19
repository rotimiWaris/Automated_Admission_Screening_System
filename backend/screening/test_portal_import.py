import base64
from io import BytesIO
from datetime import date
from django.test import TestCase
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from result_ocr import OCRError
from .table_import import import_students
from .models import RuleSet, Verification
from .test_routes import EXAM

SCHOOL = ['#','Jamb reg','Form Number','Full Name','E-mail','Phone','DOB','Sitting(s)','OLevel Result Sitting 1','OLevel Result Sitting 2','Choice 1','Jamb Programme','Jamb Subject','Jamb Score','ND Institution Name','ND Graduation Year','ND Course Of Study','ND Grade','ND CGPA','Submit Date','Remark','Unneeded']
JAMB = ['S/N','REG NO.','Name','State','UTME Subject','JAMB','Has 5 Credits?']


def school(reg='REG001', serial=1):
    return {'#':serial,'Jamb reg':reg,'Form Number':'FORM'+reg,'Full Name':'Test Student EXAMPLE','DOB':'2000-01-01','Sitting(s)':1,'OLevel Result Sitting 1':EXAM,'Choice 1':'MAPOLY','Jamb Programme':'Accountancy','ND Course Of Study':'Accountancy','ND Institution Name':'Test Polytechnic','ND Grade':'Lower Credit','ND CGPA':'2.5','Remark':'wrong old remark','Jamb Score':240,'Jamb Subject':'English Language; Mathematics; Economics; Commerce'}


def jamb(reg='REG001'):
    return {'REG NO.':reg,'Name':'EXAMPLE TEST STUDENT','UTME Subject':'English Language; Mathematics; Economics; Commerce','JAMB':240,'Has 5 Credits?':True}


def workbook(schools, jambs, reverse=False):
    book=Workbook();book.remove(book.active)
    sheets=[('School',SCHOOL,schools),('JAMB',JAMB,jambs)]
    for name,headers,rows in reversed(sheets) if reverse else sheets:
        sheet=book.create_sheet(name);sheet.append(headers)
        for row in rows:sheet.append([row.get(h,'') for h in headers])
    stream=BytesIO();book.save(stream);return stream.getvalue()


class PortalTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('load_compendium',verbosity=0)
        cls.user=get_user_model().objects.create_user(username='portal-user')

    def setUp(self):self.client.force_login(self.user)

    def process(self, raw, route='ND Full-Time', save=False):
        response=self.client.post('/api/ocr/',{'file':SimpleUploadedFile('portal.xlsx',raw),'admission_type':route})
        self.assertEqual(response.status_code,200,response.content[:300])
        imported=response.json()
        rule=RuleSet.objects.get(programme='Accountancy',active=True)
        response=self.client.post('/api/screen-batch/',{'pages':[imported['token']],'admission_type':route,'course_rules':{'Accountancy':{'id':rule.id,'revision':rule.revision}},'save':save},content_type='application/json')
        self.assertEqual(response.status_code,200,response.content[:300])
        return imported,response.json()

    def test_join_order_names_status_and_clean_export_all_routes(self):
        for route in ['ND Full-Time']:
            with self.subTest(route=route):
                raw=workbook([school('A',7),school('B',2)], [jamb('B')|{'Name':'DIFFERENT PERSON','Has 5 Credits?':False},jamb('A')],True)
                imported,result=self.process(raw,route)
                self.assertEqual([r['registration_number'] for r in result['results']],['A','B'])
                self.assertEqual([r['status'] for r in result['results']],['Eligible','Pending Review'])
                sheet=load_workbook(BytesIO(base64.b64decode(result['excel']['base64']))).active
                self.assertEqual([sheet.cell(r,1).value for r in [2,3]],['7','2'])
                self.assertEqual(sheet.cell(2,sheet.max_column).value,'Qualified')
                self.assertEqual(sheet.cell(3,sheet.max_column).value,'Pending Review')
                note=sheet.cell(3,sheet.max_column).comment.text
                self.assertIn('School name',note);self.assertIn('Has 5 Credits?',note)
                self.assertIn('Previous remark: wrong old remark',note)
                self.assertNotIn('Unneeded',[c.value for c in sheet[1]])
                draft=imported['drafts'][0]
                self.assertEqual(draft['source_record'].get('Surname'), 'EXAMPLE')
                if route=='ND Full-Time':self.assertEqual(draft['jamb_score'],240)

    def test_unmatched_and_duplicate_registrations_are_reported_not_guessed(self):
        raw=workbook([school('A'),school('MISSING'),school('DUP')],[jamb('A'),jamb('DUP'),jamb('DUP'),jamb('EXTRA')])
        drafts,_=import_students(raw,'portal.xlsx')
        report=drafts[0]['merge_report']
        self.assertEqual(report['matched'],1)
        self.assertEqual(len(report['excluded_school_rows']),2)
        self.assertEqual(report['unused_jamb_rows'],3)
        self.assertEqual(len(drafts),1)
        with self.assertRaisesMessage(OCRError,'No unambiguous registration matches'):
            import_students(workbook([school('Z')],[jamb('A')]),'portal.xlsx')

    def test_normalized_registration_and_ambiguous_name_review(self):
        drafts,_=import_students(workbook([school('reg-001')|{'Full Name':'ALL UPPER NAME'}],[jamb('REG001')]),'portal.xlsx')
        self.assertEqual(len(drafts),1)
        self.assertTrue(any('uppercase surname' in issue for issue in drafts[0]['review_issues']))

    def test_10000_students_screen_save_export_and_merged_reimport(self):
        raw=workbook([school(f'R{i:05}',i+1) for i in range(10000)],[jamb(f'R{i:05}') for i in reversed(range(10000))])
        imported,result=self.process(raw,save=True)
        self.assertEqual(result['summary'],{'Eligible':10000})
        self.assertEqual(Verification.objects.count(),10000)
        output=base64.b64decode(result['excel']['base64'])
        sheet=load_workbook(BytesIO(output),read_only=True).active
        self.assertEqual(sheet.max_row,10001)
        # Export remains a supported merged register (the workbook has long comment XML).
        drafts,_=import_students(output,'merged.xlsx')
        self.assertEqual(len(drafts),10000)
        self.assertEqual(drafts[-1]['registration_number'],'R09999')

    def test_titles_and_blank_rows_before_headers(self):
        book=load_workbook(BytesIO(workbook([school()],[jamb()])))
        for sheet in book:
            sheet.insert_rows(1,3)
            sheet['A1']='PROGRAMME:Civil Engineering Technology'
            sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=sheet.max_column)
            sheet['A2']='Admission register'
        raw=BytesIO();book.save(raw)
        drafts,_=import_students(raw.getvalue(),'titles.xlsx')
        self.assertEqual(len(drafts),1)
        self.assertEqual(drafts[0]['registration_number'],'REG001')
        self.assertEqual(drafts[0]['source_rows'],['School, row 5','JAMB, row 5'])
        self.assertEqual(drafts[0]['source_record']['#'],'1')

    def test_header_search_rejects_unlabelled_data_and_extra_values(self):
        book=Workbook();sheet=book.active
        sheet.append(['PROGRAMME:Civil Engineering Technology'])
        sheet.append(['Some person','REG001',240])
        raw=BytesIO();book.save(raw)
        with self.assertRaisesMessage(OCRError,'no recognised column headers'):
            import_students(raw.getvalue(),'invalid.xlsx')
        book=load_workbook(BytesIO(workbook([school()],[jamb()])))
        book['JAMB'].cell(2,len(JAMB)+1,'unlabelled value')
        raw=BytesIO();book.save(raw)
        with self.assertRaisesMessage(OCRError,'values exist without column headers'):
            import_students(raw.getvalue(),'invalid.xlsx')

    def test_school_utme_is_authoritative_even_when_jamb_disagrees(self):
        drafts,_=import_students(workbook([school()],[jamb()|{'JAMB':99,'UTME Subject':''}]),'portal.xlsx')
        self.assertEqual(drafts[0]['jamb_score'],240)
        self.assertEqual(len(drafts[0]['jamb_subjects']),4)
        drafts,_=import_students(workbook([school()|{'Jamb Score':'','Jamb Subject':''}],[jamb()]),'portal.xlsx')
        self.assertIsNone(drafts[0]['jamb_score'])
        self.assertEqual(drafts[0]['jamb_subjects'],[])

    def test_programme_variant_matches_without_explicit_mapping(self):
        from django.core import signing
        from .programmes import programme_key
        for label in ['Civil Engineering Technology','Civil Engineering Full Time ND']:
            self.assertEqual(programme_key(label),programme_key('Civil Engineering'))
        self.assertNotEqual(programme_key('Computer Engineering Technology'),programme_key('Computer Science'))
        raw=workbook([school()|{'Jamb Programme':'Civil Engineering Technology'}],[jamb()])
        imported=self.client.post('/api/ocr/',{'file':SimpleUploadedFile('portal.xlsx',raw)}).json()
        response=self.client.post('/api/screen-batch/',{'pages':[imported['token']],'course_rules':{},'save':False},content_type='application/json')
        self.assertEqual(response.status_code,200)
        result=response.json()['results'][0]['result']
        self.assertIsNotNone(result)
        self.assertEqual(result['rules']['programme'],'Civil Engineering')
