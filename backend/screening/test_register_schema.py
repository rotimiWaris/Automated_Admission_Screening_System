from django.test import SimpleTestCase
from result_ocr import OCRError
from result_ocr.tables import TableRow
from .register_schema import prepare_register_row
from .applicant_data import HEADERS
from .table_import import import_students
from django.conf import settings
from pathlib import Path
from io import BytesIO
from openpyxl import load_workbook

class GeneratedColumnTests(SimpleTestCase):
    def test_nd_full_time_without_generated_columns(self):
        p=Path(settings.BASE_DIR).parent/'samples/synthetic-admissions/tabular/mapoly-student-register-sample.xlsx'
        w=load_workbook(p);s=w.active
        omitted=['#','remark','age','sitting(s)','olevel result sitting 2']
        # Keep both sitting cells for the two-sitting applicant.
        omitted.remove('olevel result sitting 2')
        for index in range(s.max_column,0,-1):
            if s.cell(1,index).value in omitted:s.delete_cols(index)
        raw=BytesIO();w.save(raw)
        drafts,_=import_students(raw.getvalue(),'without-generated.xlsx')
        self.assertEqual(len(drafts),12)
        self.assertEqual(drafts[0]['original_remark'],'')
        self.assertIsInstance(drafts[0]['applicant']['age'],int)
        self.assertEqual(drafts[8]['source_record']['sitting(s)'],'2')
        self.assertEqual(drafts[0]['source_record']['#'],'1')

    def test_error_names_only_missing_required_fields(self):
        values=dict.fromkeys(HEADERS,'')
        del values['jamb score'];del values['remark'];del values['age']
        with self.assertRaises(OCRError) as caught:
            prepare_register_row(TableRow('Applicants',2,values),HEADERS,'ND Full-Time',1)
        self.assertIn('Missing required columns: jamb score.',str(caught.exception))
        self.assertNotIn('Missing required columns: remark',str(caught.exception))

    def test_duplicate_normalized_heading_explained(self):
        values=dict.fromkeys(HEADERS,'');values['REMARK']=''
        with self.assertRaisesMessage(OCRError,'appears more than once'):
            prepare_register_row(TableRow('Applicants',2,values),HEADERS,'ND Full-Time',1)
