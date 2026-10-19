"""Excel reports generated from server-evaluated batch outcomes."""
from io import BytesIO
import re
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment
from .applicant_data import HEADERS, remark_comparison
from .domain import GRADES
from .result_parser import strip_assessment_marks
from .jamb_subjects import parse_subjects

MARKS = {'passed': '✓', 'failed': '×', 'pending': '?'}


def enrich_result(result, draft, admission_type):
    result['admission_type'] = admission_type
    result['applicant'] = draft.get('applicant', {})
    result['source_record'] = draft.get('source_record', {})
    result['original_remark'] = draft.get('original_remark', '')
    result['remark'] = {'Eligible': 'Qualified', 'Not Eligible': 'Unqualified'}.get(result['status'], 'Pending Review')
    result['remark_check'] = remark_comparison(result['original_remark'], result['status'])
    best = {}
    for sitting in result['sittings']:
        for row in sitting['subjects']:
            previous = best.get(row['subject'])
            if previous is None or (row['grade'] in GRADES and (previous not in GRADES or GRADES.index(row['grade']) < GRADES.index(previous))):
                best[row['subject']] = row['grade']
    checks = []
    def add(subject, threshold, role):
        grade = best.get(subject)
        permitted = all(c['status'] == 'passed' for c in result['checks'] if c['label'] == 'Examination sittings')
        state = 'pending' if not permitted or (grade and grade not in GRADES) or (grade is None and not result['confirmed']) else 'passed' if grade in GRADES and GRADES.index(grade) <= GRADES.index(threshold) else 'failed'
        checks.append({'subject': subject, 'grade': grade or 'Missing', 'minimum_grade': threshold, 'role': role, 'status': state, 'mark': MARKS[state]})
    for required in result['rules']['required_subjects']:
        add(required['subject'], required['minimum_grade'], 'Compulsory subject')
    criteria = result['rules'].get('criteria', {})
    for path_index, path in enumerate(criteria.get('credit_paths', []), 1):
        for group_index, group in enumerate(path, 1):
            for option in group['options']:
                for subject in option:
                    add(subject, 'C6', f'Path {path_index}, group {group_index}: choose {group["count"]}; option {" / ".join(option)}')
    for group in criteria.get('pass_groups', []):
        for option in group['options']:
            for subject in option:
                add(subject, 'E8', 'Additional pass: ' + ' / '.join(option))
    result['subject_checks'] = checks
    return result


def qualifying_subjects(result):
    """Only the subjects and sittings that support a completed qualification."""
    from .criteria import allocation
    if result.get('status') != 'Eligible': return {}
    best={}
    for sitting in result['sittings']:
        for row in sitting['subjects']:
            grade=row['grade']; subject=row['subject']
            if grade in GRADES and (subject not in best or GRADES.index(grade)<GRADES.index(best[subject])):
                best[subject]=grade
    rules=result['rules']; criteria=rules.get('criteria',{})
    selected={}
    credits={s for s,g in best.items() if g in GRADES[:6]}
    for path in criteria.get('credit_paths',[]):
        chosen=allocation(path,credits)
        if chosen is not None:
            selected.update({s:'C6' for s in chosen});break
    for item in rules['required_subjects']:
        selected[item['subject']]=item['minimum_grade']
    for subject in sorted(credits):
        if len(set(selected)&credits)>=rules['minimum_credits']: break
        selected.setdefault(subject,'C6')
    for group in criteria.get('pass_groups',[]):
        chosen=allocation([group],{s for s,g in best.items() if g in GRADES[:8] and s in selected})
        if chosen is None: chosen=allocation([group],{s for s,g in best.items() if g in GRADES[:8]})
        for subject in chosen or []: selected.setdefault(subject,'E8')
    positions={}
    for subject,minimum in selected.items():
        for index,sitting in enumerate(result['sittings'],1):
            if any(r['subject']==subject and r['grade']==best.get(subject) for r in sitting['subjects']):
                positions[(index,subject)]=minimum;break
    return positions


def sized_comment(lines):
    import textwrap
    text='\n'.join(textwrap.fill(str(line),width=80,replace_whitespace=False) for line in lines)
    note=Comment(text,'Screening system')
    note.width=680
    note.height=max(160,40+18*len(text.splitlines()))
    return note


def workbook_bytes(items, admission_type):
    from .routes import headers_for
    from .mapoly_import import normal
    from .result_parser import subject_rows
    headers=headers_for(admission_type)
    display_headers = ['#', 'FORM NO', 'Jamb reg', 'NAME ON JAMB', 'SURNAME', 'FIRST NAME', 'OTHER NAME', 'E-mail', 'Phone', 'DOB', 'AGE', 'First choice', 'Jamb Programme', 'Jamb Subject', 'Jamb Score', 'Sitting(s)', 'OLevel Result Sitting 1', 'OLevel Result Sitting 2', 'JAMB OLEVEL STATUS', 'Submit Date', 'Remark'] if admission_type == 'ND Full-Time' else headers
    wb=Workbook();sheet=wb.active;sheet.title='Screenings';sheet.append(display_headers)
    def key(label):
        return next((h for h in headers if normal(h)==normal(label)),label)
    for index,item in enumerate(items,1):
        result=item.get('result') or {}
        row=dict(item.get('source_record') or result.get('source_record',{}))
        applicant=item.get('applicant') or result.get('applicant',{})
        row.setdefault('#',index)
        reg=key('JAMB Number') if admission_type.startswith('HND') else key('Jamb reg')
        row.setdefault(reg,item['registration_number'])
        row.setdefault(key('Name on Jamb'),item['student_name'])
        row.setdefault(key('jamb programme'),item['course'])
        row.setdefault(key('Department'),item['course'])
        row.setdefault(key('ND Course Of Study'),item['course'])
        row.setdefault(key('jamb score'),result.get('jamb_score'))
        row.setdefault(key('sitting(s)'),len(result.get('sittings',[])) or '')
        chosen=qualifying_subjects(result)
        used=set()
        for number in [1,2]:
            column=key(f'olevel result sitting {number}')
            exams=result.get('sittings',[])
            if number<=len(exams):
                exam=exams[number-1]
                row.setdefault(column,f'{exam["board"]} {exam["exam_type"]} {exam["year"]}; '+ '; '.join(f'{r["subject"]}: {r["grade"]}' for r in exam['subjects']))
            def annotate(match):
                entry=strip_assessment_marks(match.group(0))
                parsed=subject_rows(entry)
                if len(parsed)==1:
                    subject=parsed[0]['subject'];grade=parsed[0]['grade'];position=(number,subject)
                    if position in chosen and position not in used and grade in GRADES and GRADES.index(grade)<=GRADES.index(chosen[position]):
                        used.add(position);return entry.rstrip()+' ['+chr(0x2713)+']'
                return entry
            if row.get(column): row[column]=re.sub(r'[^;|\n\r]+',annotate,str(row[column]))
        jamb=key('jamb subject')
        if jamb in headers:
            subjects=result.get('jamb_subjects',[])
            originals=[strip_assessment_marks(part).strip() for part in re.split(r'[,;|\n]+',str(row.get(jamb,''))) if part.strip()]
            try: same=parse_subjects('; '.join(originals))==subjects
            except ValueError: same=False
            labels=originals if same or not subjects else subjects
            row[jamb]='; '.join(label+(' ['+chr(0x2713)+']' if result.get('status')=='Eligible' and subjects else '') for label in labels)
        selection = result.get('selected_nd_entry', {})
        if admission_type.startswith('HND'):
            for field in ['ND Course Of Study', 'ND Grade', 'ND CGPA', 'ND Graduation Year']:
                parts = str(row.get(field, '')).splitlines()
                position = selection.get('positions', {}).get(field)
                parts = [strip_assessment_marks(part) for part in parts]
                if result.get('status') == 'Eligible' and position is not None and position < len(parts):
                    i = position
                    parts[i] = parts[i].rstrip() + ' [' + chr(0x2713) + ']'
                row[field] = '\n'.join(parts)
        if 'age' in headers: row['age']=applicant.get('age','')
        remark=key('Remark');original=item.get('original_remark',result.get('original_remark',''))
        row[remark]={'Eligible':'Qualified','Not Eligible':'Unqualified'}.get(result.get('status'),'Pending Review')
        sheet.append([row.get(h,'') for h in headers])
        details=['Decision: '+result.get('remark','Pending Review')]
        reasons=[c['label']+': '+c['detail'] for c in result.get('checks',[]) if c['status']!='passed']
        reasons+=item.get('issues',[])
        details+=list(dict.fromkeys(reasons))
        if row[remark]=='Pending Review':details.append('Review required before a final remark can be assigned.')
        details+=[w for w in dict.fromkeys(result.get('warnings',[])+applicant.get('warnings',[])) if w not in reasons]
        if str(original).strip().casefold() not in ['', 'none', 'n/a', '-']:
            details.append('Previous remark: '+str(original))
            comparison = remark_comparison(original,result.get('status'))['detail']
            if comparison: details.append(comparison)
        details.append(admission_type)
        sheet.cell(sheet.max_row,headers.index(remark)+1).comment=sized_comment(details)
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:
        cell.fill=PatternFill('solid',fgColor='175B3A');cell.font=Font(color='FFFFFF',bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value,str):cell.data_type='s'
            cell.alignment=Alignment(vertical='top',wrap_text=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width=55 if 'olevel' in str(column[0].value).lower() else 24
    stream=BytesIO();wb.save(stream);return stream.getvalue()
