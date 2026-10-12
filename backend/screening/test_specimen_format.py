"""Fictional fixtures reproduce the supplied workbook format without applicant data."""
from io import BytesIO
from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook
from .applicant_data import HEADERS
from .table_import import import_students
from .result_parser import subject_rows, sitting
from .jamb_subjects import parse_subjects
from .models import RuleSet
from .domain import evaluate
from .screening_export import enrich_result, workbook_bytes

EXAM = """Exam Name:WAEC
Exam Type:WASSCE FOR SCHOOL CANDIDATES 2026
Exam Year:2026
Exam No:0000000000
Exam Center:EXTERNAL MODEL SCHOOL 1999
Candidate Name:FICTIONAL APPLICANT

COMMERCE - [C4] [?]
ACCOUNTING - [B3] [?]
MARKETING - [A1] [?]
ECONOMICS - [B3] [?]
GOVERNMENT - [C5] [?]
CIVIC EDUCATION - [B3] [?]
ENGLISH LANGUAGE - [C6] [?]
GENERAL MATHEMATICS - [B3] [?]
BIOLOGY - [OUTSTANDING] [?]"""

class SpecimenFormatTests(SimpleTestCase):
    def test_optional_outstanding_does_not_block_marketing_and_is_in_comment(self):
        from .test_compendium import CompendiumTests
        draft=self.draft()
        rule=CompendiumTests().rule(3)
        result=enrich_result(evaluate(rule,{**draft,'confirmed':not draft['review_issues']}),draft,'ND Full-Time')
        self.assertEqual(result['status'],'Eligible')
        self.assertTrue(any('Biology: Outstanding' in w for w in result['warnings']))
        item={**draft,'issues':result['warnings'],'result':result,'status':result['status']}
        sheet=load_workbook(BytesIO(workbook_bytes([item],'ND Full-Time'))).active
        self.assertEqual(sheet['U2'].value,'Qualified')
        self.assertIn('Biology: Outstanding',sheet['U2'].comment.text)

    def test_compulsory_outstanding_still_needs_review(self):
        from .test_compendium import CompendiumTests
        draft=self.draft(EXAM.replace('ENGLISH LANGUAGE - [C6]', 'ENGLISH LANGUAGE - [OUTSTANDING]'))
        result=evaluate(CompendiumTests().rule(3),{**draft,'confirmed':True})
        self.assertEqual(result['status'],'Pending Review')

    def test_required_additional_pass_still_needs_review(self):
        from .test_compendium import CompendiumTests
        helper=CompendiumTests()
        data=helper.data('E,M,C,B,Ec','E,B,C,Ec')
        data['sittings'][0]['subjects'].append({'subject':'Physics','grade':'Outstanding'})
        self.assertEqual(evaluate(helper.rule(10),data)['status'],'Pending Review')

    def test_unconfirmed_extraction_still_needs_review(self):
        from .test_compendium import CompendiumTests
        draft=self.draft()
        self.assertEqual(evaluate(CompendiumTests().rule(3),{**draft,'confirmed':False})['status'],'Pending Review')

    def draft(self, text=EXAM):
        row = dict.fromkeys(HEADERS, '')
        row.update({'Form No.':'FICTIONAL-FORM', 'Jamb reg':'FICTIONAL-JAMB',
                    'Surname':'Fictional', 'First Name':'Applicant', 'dob':'2005-01-01',
                    'jamb programme':'Marketing', 'jamb score':'211', 'sitting(s)':'1',
                    'jamb subject':'Economics - [55]\nMathematics - [45]\nCommerce - [52]\nEnglish Language - [211]',
                    'olevel result sitting 1':text, 'remark':'QUALIFIED'})
        w=Workbook(); w.active.append(HEADERS); w.active.append([row[h] for h in HEADERS])
        stream=BytesIO(); w.save(stream)
        drafts,_=import_students(stream.getvalue(),'specimen.xlsx')
        return drafts[0]

    def test_metadata_brackets_and_outstanding(self):
        draft=self.draft()
        exam=draft['sittings'][0]
        self.assertEqual((exam['board'],exam['exam_type'],exam['year']),('WAEC','School',2026))
        self.assertEqual(len(exam['subjects']),9)
        grades={r['subject']:r['grade'] for r in exam['subjects']}
        self.assertEqual(grades['Mathematics'],'B3')
        self.assertEqual(grades['Financial Accounting'],'B3')
        self.assertEqual(grades['Biology'],'Outstanding')
        self.assertEqual(draft['review_issues'],[])
        self.assertEqual(len(draft['jamb_subjects']),4)
        self.assertEqual(draft['jamb_score'],211)

    def test_empty_placeholder_and_completed_grades(self):
        draft=self.draft(EXAM.replace('BIOLOGY - [OUTSTANDING] [?]', '- []'))
        self.assertEqual(len(draft['sittings'][0]['subjects']),8)
        self.assertEqual(draft['review_issues'],[])

    def test_grade_statuses_and_invalid_codes(self):
        for source,expected in [('AR','Awaiting'),('ABS','Absent'),('CANCELLED','Cancelled'),('A2','Unreadable'),('A3','Unreadable'),('P7','Unreadable'),('P8','Unreadable')]:
            with self.subTest(source=source):
                self.assertEqual(subject_rows(f'BIOLOGY - [{source}] [?]')[0]['grade'],expected)
        self.assertEqual(subject_rows('BIOLOGY - [B3] [\u2713]')[0]['grade'],'B3')
        self.assertEqual(subject_rows('BIOLOGY - [F9] [\u00d7]')[0]['grade'],'F9')

    def test_jamb_abbreviations(self):
        self.assertEqual(parse_subjects('English Language - [190]\nLit. in English - [40]\nChristian Rel. Know - [50]\nPrinc. of Account - [50]'),
                         ['English Language','Literature in English','Christian Religious Studies','Financial Accounting'])

    def test_private_and_typo_metadata(self):
        for board,kind,expected in [('NABTEB','November/December, 2025','Private/GCE'),('NECO','Intenal','School'),('WAEC','WASSCE FOR PRIVATE CANDIDATES 2025','Private/GCE')]:
            self.assertEqual(sitting(f'Exam Name:{board}\nExam Type:{kind}\nExam Year:2025')['exam_type'],expected)

    def test_export_replaces_old_marks_and_preserves_source(self):
        draft=self.draft(EXAM.replace('OUTSTANDING','B3'))
        rule=RuleSet(programme='Marketing',minimum_jamb=100,required_subjects=[{'subject':'English Language','minimum_grade':'C6'},{'subject':'Mathematics','minimum_grade':'C6'}])
        result=enrich_result(evaluate(rule,{**draft,'confirmed':True}),draft,'ND Full-Time')
        item={**draft,'issues':[],'result':result,'status':result['status']}
        sheet=load_workbook(BytesIO(workbook_bytes([item],'ND Full-Time'))).active
        self.assertEqual(sheet.max_column,21)
        self.assertIn('GENERAL MATHEMATICS - [B3] [\u2713]',sheet['Q2'].value)
        self.assertNotIn('[?] [',sheet['Q2'].value)
        self.assertIn('English Language - [211]',sheet['N2'].value)
        self.assertIn('Exam No:0000000000',sheet['Q2'].value)
        self.assertEqual(sheet['U2'].value,'Qualified')
