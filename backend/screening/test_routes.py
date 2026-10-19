import base64
import json
from io import BytesIO
from django.test import TestCase
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from django.contrib.auth import get_user_model
from openpyxl import Workbook, load_workbook
from .routes import headers_for
from .models import RuleSet, Verification
from .identities import candidate_name_issues

EXAM="""Exam Name:WAEC
Exam Type:WASSCE FOR SCHOOL CANDIDATES 2025
Exam Year:2025
Candidate Name:EXAMPLE TEST STUDENT
ENGLISH LANGUAGE - [C6] [?]
GENERAL MATHEMATICS - [B3] [?]
ECONOMICS - [B3] [?]
COMMERCE - [C4] [?]
ACCOUNTING - [B3] [?]
GOVERNMENT - [A1] [?]
Unfamiliar elective - [A1] [?]"""

class RouteTests(TestCase):
    def setUp(self):
        from django.contrib.auth import get_user_model
        self.client.force_login(get_user_model().objects.create_user(username='authenticated-fixture'))

    @classmethod
    def setUpTestData(cls): call_command('load_compendium',verbosity=0)
    def run_route(self,route,changes=None,save=False,omit=(),smart=False):
        headers=[h for h in headers_for(route) if h not in omit];values=dict.fromkeys(headers,'')
        values.update({'#':'1','Surname':'Example','First name':'Test','Other name':'Student','Sitting(s)':'1','OLevel Result Sitting 1':EXAM,'Remark':''})
        if route.startswith('HND'):
            values.update({'JAMB Number':'ROUTE-001','Form Number':'FORM-001','First Choice':'Preferred institution','ND Course Of Study':'Accountancy','ND Institution Name':'Fictional Polytechnic','ND Grade':'Lower Credit','ND CGPA':'2.50','ND Graduation Year':'2024'})
        else: values.update({'jamb reg':'ROUTE-001','Form number':'FORM-001','Department':'Accountancy','Date of Birth':'2004-01-01'})
        values.update(changes or {})
        w=Workbook();w.active.append(headers);w.active.append([values[h] for h in headers]);raw=BytesIO();w.save(raw)
        response=self.client.post('/api/ocr/',{'file':SimpleUploadedFile('routes.xlsx',raw.getvalue()),'admission_type':route,'smart_processing':str(smart).lower()})
        self.assertEqual(response.status_code,200,response.content[:300]);token=response.json()['token']
        rule=RuleSet.objects.get(programme='Accountancy',active=True)
        response=self.client.post('/api/screen-batch/',json.dumps({'pages':[token],'admission_type':route,'course_rules':{'Accountancy':{'id':rule.id,'revision':rule.revision}},'save':save}),content_type='application/json')
        self.assertEqual(response.status_code,200,response.content[:300])
        data=response.json();sheet=load_workbook(BytesIO(base64.b64decode(data['excel']['base64']))).active
        return data['results'][0],sheet

    def test_route_layouts_academic_only_and_minimal_ticks(self):
        for route in ['ND Part-Time','HND Full-Time','HND Part-Time']:
            with self.subTest(route=route):
                item,sheet=self.run_route(route)
                self.assertEqual(item['status'],'Eligible')
                self.assertEqual(sheet.max_column,len(headers_for(route)))
                self.assertEqual(sheet.cell(2,sheet.max_column).value,'Qualified')
                self.assertNotIn('Previous',sheet.cell(2,sheet.max_column).comment.text)
                self.assertFalse(any('JAMB' in c['label'] for c in item['result']['checks']))
                text=sheet.cell(2,headers_for(route).index('OLevel Result Sitting 1')+1).value
                self.assertEqual(text.count(chr(0x2713)),5)
                self.assertIn('Unfamiliar elective - [A1]',text)
                self.assertNotIn('[?]',text)

    def test_name_mismatch_pending_with_explanation(self):
        item,sheet=self.run_route('ND Part-Time',{'Other name':'Different'})
        self.assertEqual(item['status'],'Pending Review')
        self.assertIn('does not match',sheet.cell(2,15).comment.text)
        self.assertIn('Different',sheet.cell(2,15).comment.text)
        self.assertEqual(sheet.cell(2,15).value,'Pending Review')
        self.assertEqual(candidate_name_issues('Test Example Student','Candidate Name:EXAMPLE TEST STUDENT'),[])
        self.assertTrue(candidate_name_issues('Test Example','Candidate Name:EXAMPLE TEST STUDENT'))

    def test_hnd_grade_boundaries_conflicts_and_missing_values(self):
        for cgpa,grade,status in [('1.99','Fail','Not Eligible'),('2.00','Pass','Pending Review'),('2.49','Pass','Pending Review'),('2.50','Lower Credit','Eligible'),('2.99','Lower Credit','Eligible'),('3.00','Upper Credit','Eligible'),('3.50','Distinction','Eligible'),('4.00','Distinction','Eligible'),('4.01','Distinction','Pending Review'),('2.50','Upper Credit','Not Eligible'),('','Lower Credit','Pending Review')]:
            with self.subTest(cgpa=cgpa,grade=grade):
                item,sheet=self.run_route('HND Full-Time',{'ND CGPA':cgpa,'ND Grade':grade})
                self.assertEqual(item['status'],status)

    def test_unknown_subject_when_credits_insufficient_is_reviewed(self):
        exam=EXAM.replace('ECONOMICS - [B3] [?]','Unknown economics - [B3]').replace('COMMERCE - [C4] [?]','Unknown commerce - [C4]')
        item,sheet=self.run_route('ND Part-Time',{'OLevel Result Sitting 1':exam})
        self.assertEqual(item['status'],'Pending Review')
        self.assertIn('unrecognised entry',sheet.cell(2,15).comment.text)

    def test_route_saves_and_repeat_export_does_not_duplicate_history(self):
        user=get_user_model().objects.create_user(username='route-user');self.client.force_login(user)
        self.run_route('HND Part-Time',save=True)
        item,sheet=self.run_route('HND Part-Time',save=True)
        self.assertEqual(item['status'],'Eligible')
        self.assertTrue(item['save_skipped'])
        self.assertEqual(sheet.cell(2,19).value,'Qualified')
        self.assertEqual(Verification.objects.count(),1)

    def test_generated_columns_can_be_omitted_in_all_new_routes(self):
        for route in ['ND Part-Time','HND Full-Time','HND Part-Time']:
            with self.subTest(route=route):
                item,sheet=self.run_route(route,omit=['#','Remark','Sitting(s)','OLevel Result Sitting 2'])
                self.assertEqual(item['status'],'Eligible')
                self.assertEqual(sheet.max_column,len(headers_for(route)))
                self.assertEqual(sheet.cell(2,sheet.max_column).value,'Qualified')
                self.assertNotIn('Previous',sheet.cell(2,sheet.max_column).comment.text)
                self.assertEqual(str(sheet.cell(2,1).value),'1')

    def test_smart_name_notice_survives_screening_and_export(self):
        for route in ['ND Part-Time','HND Full-Time','HND Part-Time']:
            with self.subTest(route=route):
                strict,_=self.run_route(route,{'Surname':'Exampel'})
                self.assertEqual(strict['status'],'Pending Review')
                item,sheet=self.run_route(route,{'Surname':'Exampel'},smart=True)
                self.assertEqual(item['status'],'Eligible')
                self.assertTrue(item['result']['smart_processing'])
                self.assertTrue(any('minor name typo' in warning for warning in item['result']['warnings']))
                note=sheet.cell(2,sheet.max_column).comment.text
                self.assertIn('minor name typo',note)
                self.assertIn('Exampel',note)

    def test_other_routes_ignore_jamb_sheet_and_extra_portal_fields(self):
        from .test_portal_import import workbook, school, jamb
        from .table_import import import_students
        for route in ['ND Part-Time','HND Full-Time','HND Part-Time']:
            with self.subTest(route=route):
                raw=workbook([school('SCHOOLONLY')],[jamb('DIFFERENT')|{'Has 5 Credits?':False}])
                drafts,_=import_students(raw,'portal.xlsx',route)
                self.assertEqual(len(drafts),1)
                self.assertEqual(drafts[0]['registration_number'],'FORMSCHOOLONLY')
                self.assertEqual(set(drafts[0]['source_record']),set(headers_for(route)))
                self.assertNotIn('jamb score',drafts[0]['source_record'])
                self.assertFalse(any('Has 5 Credits' in issue for issue in drafts[0]['review_issues']))
                from .routes import evaluate_route
                rule=RuleSet.objects.get(programme='Accountancy',active=True)
                result=evaluate_route(rule,drafts[0],route,True)
                self.assertEqual(result['status'],'Eligible')
                self.assertFalse(any(c['label']=='JAMB score' for c in result['checks']))
                if route.startswith('HND'):
                    self.assertTrue(any(c['label']=='ND CGPA' for c in result['checks']))
                    self.assertNotIn('Admission age',[c['label'] for c in result['checks']])

    def test_hnd_aligned_entries_select_qualifying_credentials_and_comments(self):
        for route in ['HND Full-Time', 'HND Part-Time']:
            item, sheet = self.run_route(route, {'ND Course Of Study': 'Unknown profession\nAccountnacy', 'ND Grade': 'Distinction\nLower Credit', 'ND CGPA': '3.80\n2.70', 'JAMB Number': ''})
            self.assertEqual(item['status'], 'Eligible')
            self.assertEqual(item['result']['selected_nd_entry']['index'], 1)
            for field in ['ND Course Of Study','ND Grade','ND CGPA']:
                value = sheet.cell(2, headers_for(route).index(field)+1).value
                self.assertNotIn(chr(0x2713), value.splitlines()[0])
                self.assertIn(chr(0x2713), value.splitlines()[1])
            self.assertIn('selected entry 2', sheet.cell(2,19).comment.text)

    def test_course_typos_repetition_and_unrelated_courses(self):
        for route in ['ND Part-Time','HND Full-Time','HND Part-Time']:
            field = 'ND Course Of Study' if route.startswith('HND') else 'Department'
            for course, status in [('Accountnacy','Eligible'),('Accountancy Accountancy','Eligible'),('Marine Navigation','Not Eligible')]:
                item, sheet = self.run_route(route, {field: course})
                self.assertEqual(item['status'], status)
                self.assertIsNotNone(sheet.cell(2,sheet.max_column).comment)

    def test_no_jamb_column_needed_but_form_is_required(self):
        for route in ['ND Part-Time','HND Full-Time','HND Part-Time']:
            jamb = 'JAMB Number' if route.startswith('HND') else 'jamb reg'
            form = 'Form Number' if route.startswith('HND') else 'Form number'
            item, _ = self.run_route(route, omit=[jamb])
            self.assertEqual(item['status'],'Eligible')
            item, sheet = self.run_route(route, {form: ''})
            self.assertEqual(item['status'],'Pending Review')
            self.assertIn('Form number is required',sheet.cell(2,sheet.max_column).comment.text)

    def test_misaligned_entries_are_not_combined(self):
        item, sheet = self.run_route('HND Full-Time', {'ND Course Of Study':'Accountancy\nMarketing\nAccountancy', 'ND Grade':'Lower Credit\nUpper Credit','ND CGPA':'2.70\n3.20'})
        self.assertEqual(item['status'],'Pending Review')
        self.assertIn('different numbers of lines',sheet.cell(2,19).comment.text)

    def test_aligned_entries_do_not_borrow_a_grade_from_another_line(self):
        item, sheet = self.run_route('HND Part-Time', {'ND Course Of Study':'Accountancy\nAccountancy', 'ND Grade':'Upper Credit\nLower Credit', 'ND CGPA':'2.70\n3.20'})
        self.assertEqual(item['status'],'Not Eligible')
        self.assertIn('CGPA implies',sheet.cell(2,19).comment.text)

    def test_unrelated_course_is_saved_and_explained(self):
        item, sheet = self.run_route('HND Full-Time', {'ND Course Of Study':'Marine Navigation'},save=True)
        self.assertEqual(item['status'],'Not Eligible')
        self.assertIn('does not match',sheet.cell(2,19).comment.text)
        self.assertEqual(Verification.objects.count(),1)
